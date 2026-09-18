"""Extraction des listes d'actes spécialisés par GN (score de réadaptation SPÉCIALISÉE).

Fichier ATIH `ACTES_listes_SPE.xlsx` (annexe 4 du volume 3 du manuel de
groupage GME) : deux feuilles à joindre pour obtenir directement des paires
exploitables (GN, acte) —

- `gn_liste` : GN -> identifiant de liste (GN_liste). Une même liste est
  souvent partagée par plusieurs GN (regroupement). Valeur 'PAS DE LISTE'
  (ou vide) = ce GN n'a pas de notion de réadaptation spécialisée du tout
  (ex. 0103, 0118, 0134 — cf. section 3.4.1.2 du manuel : ces GN sont
  orientés directement sans test de score).
- `total_liste` : acte (CSARR ou CCAM) -> GN_liste auquel il appartient en
  tant qu'acte marqueur. Un acte peut apparaître dans plusieurs listes
  (une ligne par liste).

`load_from_xlsx` fait la jointure au chargement (plutôt que de stocker les
deux feuilles séparément + joindre à chaque requête) : le résultat est déjà
une paire (gn, code_acte) directement utilisable pour tester "cet acte
réalisé pendant CE séjour compte-t-il dans son score spécialisé ?".
"""
from __future__ import annotations

from pathlib import Path

import openpyxl


def load_from_xlsx(path: str | Path) -> list[dict]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)

    gn_to_liste: dict[str, str] = {}
    ws_gn = wb["gn_liste"]
    it_gn = ws_gn.iter_rows(values_only=True)
    next(it_gn)  # header
    for row in it_gn:
        if not row or not row[0]:
            continue
        gn = str(row[0]).strip()
        gn_liste = row[2] if len(row) > 2 else None
        if gn_liste and str(gn_liste).strip().upper() != "PAS DE LISTE":
            gn_to_liste[gn] = str(gn_liste).strip()

    liste_to_actes: dict[str, list[tuple[str, str]]] = {}
    ws_total = wb["total_liste"]
    it_total = ws_total.iter_rows(values_only=True)
    next(it_total)  # header
    for row in it_total:
        if not row or not row[0]:
            continue
        code_acte = str(row[0]).strip()
        type_acte = str(row[2]).strip().upper() if len(row) > 2 and row[2] else None
        gn_liste = row[5] if len(row) > 5 else None
        if not gn_liste:
            continue
        liste_to_actes.setdefault(str(gn_liste).strip(), []).append((code_acte, type_acte))
    wb.close()

    seen: set[tuple[str, str]] = set()
    rows: list[dict] = []
    for gn, gn_liste in gn_to_liste.items():
        for code_acte, nomenclature in liste_to_actes.get(gn_liste, []):
            key = (gn, code_acte)
            if key in seen:
                continue
            seen.add(key)
            rows.append({"gn": gn, "code_acte": code_acte, "nomenclature": nomenclature})
    return rows

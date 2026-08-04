"""Extraction du fichier associé CSAR (xlsx ATIH) + hiérarchie et transcodage CSARR/CSAR.

Contrairement à CIM-10, CCAM et CSARR, le CSAR n'a qu'un seul niveau de
hiérarchie : 11 chapitres, directement identifiés par les 2 premiers
caractères du code de l'acte (ex. "01E01" -> chapitre "01"). Il n'y a pas de
sous-chapitre/paragraphe, et le fichier xlsx ne contient nulle part le
libellé des chapitres eux-mêmes — seulement les codes d'actes (feuille
"actes_CSAR", colonnes Code_CSAR/Libelle_CSAR/Exclusions).

Les libellés des 11 chapitres sont donc repris tels quels du guide de
codage CSAR (guide_csar_version_n1_2026.pdf, Annexe 1 "Liste des codes" —
chaque chapitre y a son propre en-tête "CHAPITRE N : LIBELLÉ", et le §1.2
"Structure du catalogue CSAR" les liste aussi sous forme de tableau) : une
transcription ponctuelle, pas une extraction automatisée, faute d'un
fichier source structuré pour cette information.

Le fichier associé contient aussi une feuille "transcodage_CSARR_CSAR"
(table d'aide au codage, Annexe 4 du guide) : pour chaque code CSARR, le
code CSAR suggéré comme équivalent. Certaines lignes n'ont pas de
correspondance dans un sens ou l'autre (code "ZZZ") : 3 actes CSARR n'ont
pas d'équivalent CSAR, et 9 actes CSAR ont été créés sans équivalent
CSARR — la clé naturelle de cette table est donc (code_csarr, code_csar)
et non code_csarr seul, qui se répète (ZZZ) pour les 9 lignes du second cas.
"""
from __future__ import annotations

from pathlib import Path

import openpyxl

CHAPTER_LABELS = {
    "01": "FONCTIONS CEREBRALES",
    "02": "FONCTIONS SENSORIELLES ET DOULEUR",
    "03": "FONCTIONS DE LA VOIX ET DE LA PAROLE",
    "04": "FONCTIONS CARDIAQUES, VASCULAIRES ET RESPIRATOIRES",
    "05": "FONCTIONS DIGESTIVES ET NUTRITION",
    "06": "FONCTIONS GENITO-URINAIRES ET REPRODUCTIVES",
    "07": "FONCTIONS DE L'APPAREIL LOCOMOTEUR ET LIEES AU MOUVEMENT",
    "08": "FONCTIONS DE LA PEAU ET DES PHANERES",
    "09": "APPAREILLAGE",
    "10": "EDUCATION ET PREVENTION",
    "11": "ACTIVITE ET PARTICIPATION",
}


def load_flat_from_xlsx(path: str | Path) -> list[dict]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb["actes_CSAR"]
    it = ws.iter_rows(values_only=True)
    next(it)  # header
    rows = []
    for r in it:
        if not r[0]:
            continue
        code = str(r[0]).strip()
        rows.append({"code": code, "libelle": (r[1] or "").strip(), "parent_code": code[:2]})
    wb.close()
    return rows


def load_hierarchie_from_xlsx(_path: str | Path) -> list[dict]:
    return [
        {"code": code, "kind": "chapitre", "parent_code": None, "libelle": libelle}
        for code, libelle in CHAPTER_LABELS.items()
    ]


def load_transcodage_from_xlsx(path: str | Path) -> list[dict]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb["transcodage_CSARR_CSAR"]
    it = ws.iter_rows(values_only=True)
    next(it)  # header
    rows = []
    for r in it:
        if not r[0]:
            continue
        rows.append(
            {
                "code_csarr": str(r[0]).strip(),
                "libelle_csarr": (r[1] or "").strip(),
                "code_csar": str(r[2]).strip() if r[2] else None,
                "libelle_csar": (r[3] or "").strip() if r[3] else None,
            }
        )
    wb.close()
    return rows

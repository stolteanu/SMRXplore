"""Extraction des pondérations d'actes (CSARR/CCAM) pour le calcul du GR (groupage SMR).

Le fichier ATIH `ACTES_ponderations.xlsx` fournit sa propre documentation
des colonnes dans l'onglet "lisez-moi" — reprise ici telle quelle plutôt
que devinée :

- Type : D=dédié, D/ND=dédié ou non dédié, C=collectif, A=appareillage avec
  étapes, PP=pluriprofessionnel, CCAM=(ligne CCAM).
- Statut : Pond_u=pondération unique pour tous les intervenants,
  Diff_inter=différenciée selon l'intervenant, Pond_0=pondération à 0 pour
  certains couples acte/intervenant (les autres identiques), Pond_0/Diff_inter
  = combinaison des deux.
- Intervenant '00'/lib_pmsi 'tous' = la pondération s'applique à tout
  intervenant (une seule ligne pour l'acte) ; sinon une ligne par
  intervenant (jusqu'à 32 lignes pour un même acte code+libellé quand la
  pondération est différenciée — cf. statut Diff_inter).
- Hiera = subdivision CSARR (ex. '07.02.02') ; valeur spéciale 'suppr' pour
  un acte retiré du catalogue courant (le fichier couvre l'historique
  depuis 2012, donc des codes absents de nomenclature_csarr aujourd'hui).
- Validite = '×' si le code est en cours de validité, vide sinon (couplé à
  `fin`, l'année de retrait).
- mod_HW/mod_LJ/mod_XH/mod_L3 = éligibilité de l'acte à une majoration de
  pondération pour ce modulateur de lieu (valeurs des majorations dans
  l'onglet "modulateurs", feuille séparée ci-dessous).

Vérifié empiriquement : (code, intervenant) est la clé naturelle exacte —
2359 lignes, 0 doublon sur ce couple, bien que `code` seul soit dupliqué
jusqu'à 32 fois pour les actes à pondération différenciée par intervenant.

La feuille "Intervenants" du même classeur n'est PAS chargée séparément :
elle ne fait que redonner la même liste code+libellé déjà couverte par
nomenclature_csarr_intervenants (32 lignes identiques), sans information
supplémentaire.
"""
from __future__ import annotations

from pathlib import Path

import openpyxl


def load_actes_from_xlsx(path: str | Path) -> list[dict]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb["ponderations"]
    it = ws.iter_rows(values_only=True)
    next(it)  # header
    rows = []
    for r in it:
        if not r[0]:
            continue
        rows.append(
            {
                "code": str(r[0]).strip(),
                "nomenclature": r[1],
                "type": r[2],
                "statut": r[3],
                "intervenant": str(r[4]).strip() if r[4] is not None else None,
                "lib_pmsi": r[5],
                "hiera": r[6],
                "liblong": r[7],
                "ponderation_patient": r[8],
                "validite": r[9],
                "debut": r[10],
                "fin": r[11],
                "mod_hw": r[12],
                "mod_lj": r[13],
                "mod_xh": r[14],
                "mod_l3": r[15],
            }
        )
    wb.close()
    return rows


def load_modulateurs_from_xlsx(path: str | Path) -> list[dict]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb["modulateurs"]
    it = ws.iter_rows(values_only=True)
    next(it)  # header
    rows = []
    for r in it:
        if not r[0]:
            continue
        rows.append(
            {
                "code": str(r[0]).strip(),
                "libelle": r[1],
                "majoration_individuel": r[2],
                "majoration_collectif": r[3],
            }
        )
    wb.close()
    return rows

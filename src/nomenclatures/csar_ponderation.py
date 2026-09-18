"""Extraction des données de pondération CSAR depuis CSAR_infos.xlsx (ATIH),
annexe 7 du Manuel des GME vol.1/vol.3.

Contrairement au fichier associé CSAR "opérationnel" (fichier_associe_csar_*.xlsx,
déjà chargé dans nomenclature_csar/csar_hierarchie/csar_transcodage_csarr), ce
classeur documentaire est la seule source qui donne :
- le transcodage détaillé CSAR -> CSARR par couple (code_csar, intervenant,
  modalité collective) — nécessaire car la fonction de groupage ATIH calcule
  la pondération d'un acte CSAR en le transcodant d'abord en acte CSARR de
  référence (Manuel des GME vol.1, §3.1.1 et §3.3.1.3-3.3.1.4) ;
- les valeurs numériques du modulateur de temps CSAR (T0-T4) et des
  majorations du modulateur de lieu CSAR (L1/L2/L3), qui n'existent nulle
  part ailleurs dans les fichiers déjà chargés.

Feuille "transcodage_CSAR_détail" : colonnes code_csar, libelle_csar,
intervenant, acte_coll, code_csarr, libelle_csarr, commentaire. Vérifié
empiriquement (2026-09-17) : (code_csar, intervenant, acte_coll) est la clé
naturelle exacte — 5408 lignes, aucune ligne sans cible CSARR. La colonne
acte_coll vaut '0'/'1' quand la modalité collective change le transcodage
(736 couples code/intervenant avec les deux valeurs), ou '2' quand elle n'a
aucun effet sur le transcodage (le codeur peut alors coder l'acte CSAR en
individuel ou collectif, le transcodage retenu est le même dans les deux cas).

Feuille "modulateur_modalite_extension" : reprend les valeurs du Tableau 3
et du §3.3.1.4 du Manuel des GME vol.1 — pondération du modulateur de temps
(Vide=0, T0=25, T1=50, T2=80, T3=110, T4=140) et majoration du modulateur de
lieu (L1/L3 : 30 individuel / 5 collectif ; L2 : 60 individuel / sans objet
collectif — ces deux dernières colonnes ne sont renseignées que pour les
lignes 'Lieu', pas 'Temps').
"""
from __future__ import annotations

from pathlib import Path

import openpyxl


def load_transcodage_detail_from_xlsx(path: str | Path) -> list[dict]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb["transcodage_CSAR_détail"]
    it = ws.iter_rows(values_only=True)
    next(it)  # header
    rows = []
    for r in it:
        if not r[0]:
            continue
        rows.append(
            {
                "code_csar": str(r[0]).strip(),
                "intervenant": str(r[2]).strip() if r[2] is not None else None,
                "acte_coll": str(r[3]).strip() if r[3] is not None else "0",
                "code_csarr": str(r[4]).strip() if r[4] else None,
            }
        )
    wb.close()
    return rows


def load_ponderation_temps_from_xlsx(path: str | Path) -> list[dict]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb["modulateur_modalite_extension"]
    it = ws.iter_rows(values_only=True)
    next(it)  # header
    rows = []
    for r in it:
        if r[1] != "Temps":
            continue
        modalite = str(r[2]).strip() if r[2] else "Vide"
        try:
            pond = float(r[5]) if r[5] is not None else 0.0
        except (TypeError, ValueError):
            pond = 0.0
        rows.append({"modalite": modalite, "ponderation": pond})
    wb.close()
    return rows


def load_majoration_lieu_from_xlsx(path: str | Path) -> list[dict]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb["modulateur_modalite_extension"]
    it = ws.iter_rows(values_only=True)
    next(it)  # header
    rows = []
    for r in it:
        if r[1] != "Lieu":
            continue
        rows.append(
            {
                "modalite": str(r[2]).strip(),
                "majoration_individuel": r[6],
                "majoration_collectif": r[7],
            }
        )
    wb.close()
    return rows
